"""Exporta Excel de API Contexto y Base de Conocimiento a Markdown en docs/.

Uso (local o CI):
  python scripts/export_docs_from_excel.py \\
    --api-excel \"C:/path/Documentacion_API_Productos 1.xlsx\" \\
    --kb-excel \"C:/path/Base de Conocimiento IA VF01.xlsx\"

Variables de entorno (contenedores / CI):
  GENESIS_API_EXCEL_PATH
  GENESIS_KB_EXCEL_PATH
  GENESIS_DOCS_OUT_DIR   (default: docs/)
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]


def _cell(ws, r: int, c: int) -> str:
    v = ws.cell(r, c).value
    if v is None:
        return ""
    return str(v).strip()


def _md_escape_cell(text: str) -> str:
    return (text or "").replace("|", "\\|").replace("\n", " ")


def export_api_docs(excel: Path, out_dir: Path) -> list[Path]:
    wb = openpyxl.load_workbook(excel, data_only=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    # --- README ---
    ws = wb["Resumen API"] if "Resumen API" in wb.sheetnames else wb[wb.sheetnames[0]]
    props: dict[str, str] = {}
    for r in range(1, ws.max_row + 1):
        k, v = _cell(ws, r, 1), _cell(ws, r, 2)
        if k and v and k.lower() not in ("propiedad", "campo raíz", "campo raiz"):
            props[k] = v

    root_fields: list[tuple[str, str, str]] = []
    in_root = False
    for r in range(1, ws.max_row + 1):
        a, b, c = _cell(ws, r, 1), _cell(ws, r, 2), _cell(ws, r, 3)
        if a.lower() in ("campo raíz", "campo raiz"):
            in_root = True
            continue
        if in_root and a:
            root_fields.append((a, b, c))

    readme = out_dir / "README.md"
    lines = [
        "# API de Contexto — Productos por Cliente",
        "",
        "Fuente: `Documentacion_API_Productos`. Documentación canónica del servicio",
        "que alimenta el `CustomerContextSnapshot` de la banca conversacional.",
        "",
        "> En contenedores, montar o regenerar este directorio con",
        "> `GENESIS_API_EXCEL_PATH` + `scripts/export_docs_from_excel.py`.",
        "",
        "## Endpoint",
        "",
        "| Propiedad | Valor |",
        "|-----------|-------|",
    ]
    for k, v in props.items():
        if k.startswith("Retorna"):
            continue
        lines.append(f"| {_md_escape_cell(k)} | {_md_escape_cell(v)} |")
    lines += [
        "",
        "## Descripción",
        "",
        "Retorna todos los productos financieros asociados a un cliente. Incluye",
        "cuentas de ahorro, certificados de depósito, tarjetas de crédito y préstamos.",
        "",
        "## Campos raíz de la respuesta",
        "",
        "| Campo | Tipo | Descripción | Uso en conversación |",
        "|-------|------|-------------|---------------------|",
    ]
    usage = {
        "resultCode": "Validar éxito (0) antes de mapear productos",
        "resultMessage": "Log / diagnóstico; no mostrar al cliente",
        "products": "Fuente del snapshot de portafolio del turno",
    }
    for name, typ, desc in root_fields:
        lines.append(
            f"| `{name}` | {typ} | {_md_escape_cell(desc)} | {usage.get(name, '—')} |"
        )
    lines += [
        "",
        "## Índice",
        "",
        "- [Diccionario de campos](diccionario-campos.md)",
        "- [Tablas de referencia](tablas-referencia.md)",
        "- [Casos de uso por producto](casos-de-uso.md)",
        "- [Ejemplo de respuesta](ejemplo-productos.md)",
        "- [Notas operativas](notas.md)",
        "",
    ]
    readme.write_text("\n".join(lines), encoding="utf-8")
    written.append(readme)

    # --- Diccionario ---
    ws = wb["Diccionario de Campos"]
    dict_path = out_dir / "diccionario-campos.md"
    dlines = [
        "# Diccionario de campos — API Productos",
        "",
        "Relación campo → significado → uso en la banca conversacional.",
        "",
        "| Sección | Campo | Tipo | Aplica a | Descripción |",
        "|---------|-------|------|----------|-------------|",
    ]
    for r in range(2, ws.max_row + 1):
        sec, campo, typ, aplica, desc = (
            _cell(ws, r, 1),
            _cell(ws, r, 2),
            _cell(ws, r, 3),
            _cell(ws, r, 4),
            _cell(ws, r, 5),
        )
        if not campo:
            continue
        dlines.append(
            f"| {_md_escape_cell(sec)} | `{campo}` | {_md_escape_cell(typ)} | "
            f"{_md_escape_cell(aplica)} | {_md_escape_cell(desc)} |"
        )
    dlines += [
        "",
        "## Mapeo a CustomerContextSnapshot",
        "",
        "| Campo API | Snapshot | Notas |",
        "|-----------|----------|-------|",
        "| `productCategory` | `product_type` | CA→SAVINGS, CC→CHECKING, CD→TERM_DEPOSIT, TC→CREDIT_CARD, PR→LOAN |",
        "| `productIdentification` | `product_id` | Identificador único |",
        "| `productDescription` | `alias` | Nombre visible |",
        "| `productStatus` | `status` | 1/3 activos según reglas Core |",
        "| `currencyCode` | `currency` | 214→DOP, 840→USD |",
        "| `availableBalance` | `available_balance` / `credit_limit` | Semántica por categoría |",
        "| `currentBalance` | `ledger_balance` | Adeudo TC / saldo cuenta |",
        "| `maskedCardNumber` | `card_mask` | Solo TC |",
        "| `pendingBalancePr` | `LoanSnapshot.overdue_amount` | Mora, no cuota contractual |",
        "",
    ]
    dict_path.write_text("\n".join(dlines), encoding="utf-8")
    written.append(dict_path)

    # --- Tablas referencia ---
    ws = wb["Tablas de Referencia"]
    ref_path = out_dir / "tablas-referencia.md"
    rlines = [
        "# Tablas de referencia",
        "",
        "## Categorías de producto",
        "",
        "| Código | Descripción | Tipo conversacional |",
        "|--------|-------------|---------------------|",
        "| CA | Cuenta de Ahorros | SAVINGS |",
        "| CC | Cuenta Corriente | CHECKING |",
        "| CD | Certificado de Depósito | TERM_DEPOSIT |",
        "| TC | Tarjeta de Crédito | CREDIT_CARD |",
        "| PR | Préstamo | LOAN |",
        "",
        "## Estados",
        "",
        "| Código | Descripción |",
        "|--------|-------------|",
        "| 1 | Activo |",
        "| 2 | Inactivo / Bloqueado |",
        "| 3 | Activo (variante — usado en cuentas de ahorro activas) |",
        "",
        "## Monedas",
        "",
        "| Código | Moneda |",
        "|--------|--------|",
        "| 214 | DOP — Peso Dominicano |",
        "| 840 | USD — Dólar Estadounidense |",
        "",
    ]
    ref_path.write_text("\n".join(rlines), encoding="utf-8")
    written.append(ref_path)

    # --- Casos de uso ---
    ws = wb["Casos de Uso"]
    cu_path = out_dir / "casos-de-uso.md"
    clines = [
        "# Casos de uso — semántica por producto",
        "",
        "Reglas que el orquestador y los templates de respuesta deben respetar",
        "para no alucinar significados de saldo.",
        "",
        "| Tipo | Código | Regla / uso |",
        "|------|--------|-------------|",
    ]
    for r in range(2, ws.max_row + 1):
        t, c, rule = _cell(ws, r, 1), _cell(ws, r, 2), _cell(ws, r, 3)
        if not t and not rule:
            continue
        clines.append(f"| {_md_escape_cell(t)} | {_md_escape_cell(c)} | {_md_escape_cell(rule)} |")
    clines.append("")
    cu_path.write_text("\n".join(clines), encoding="utf-8")
    written.append(cu_path)

    # --- Ejemplo ---
    ws = wb["Ejemplo Productos"]
    ex_path = out_dir / "ejemplo-productos.md"
    headers = [_cell(ws, 1, c) for c in range(1, ws.max_column + 1) if _cell(ws, 1, c)]
    elines = [
        "# Ejemplo de productos",
        "",
        "Payload de muestra (valores de laboratorio). No usar en producción.",
        "",
    ]
    for r in range(2, ws.max_row + 1):
        cat = _cell(ws, 1, 1)  # wrong - use row r
        row_vals = {_cell(ws, 1, c): _cell(ws, r, c) for c in range(1, len(headers) + 1)}
        if not any(row_vals.values()):
            continue
        title = row_vals.get("productCategory") or row_vals.get("productDescription") or f"row-{r}"
        elines.append(f"## {title} — {row_vals.get('productDescription', '')}")
        elines.append("")
        elines.append("| Campo | Valor |")
        elines.append("|-------|-------|")
        for h in headers:
            if row_vals.get(h):
                elines.append(f"| `{h}` | {_md_escape_cell(row_vals[h])} |")
        elines.append("")
    ex_path.write_text("\n".join(elines), encoding="utf-8")
    written.append(ex_path)

    # --- Notas ---
    if "Notas" in wb.sheetnames:
        ws = wb["Notas"]
        npath = out_dir / "notas.md"
        nlines = ["# Notas operativas", ""]
        for r in range(1, ws.max_row + 1):
            a, b = _cell(ws, r, 1), _cell(ws, r, 2)
            if a.isdigit() and b:
                nlines.append(f"{a}. {b}")
            elif b and a.lower() not in ("notas importantes", "nº", "n°", "no"):
                nlines.append(f"- {b}")
        nlines.append("")
        npath.write_text("\n".join(nlines), encoding="utf-8")
        written.append(npath)

    return written


def export_kb_rules(excel: Path, out_dir: Path, from_row: int = 79) -> list[Path]:
    """Exporta reglas de negocio enriquecidas (routing metadata) desde Matriz_Cuentas."""
    wb = openpyxl.load_workbook(excel, data_only=True)
    ws = wb["Matriz_Cuentas"] if "Matriz_Cuentas" in wb.sheetnames else wb[wb.sheetnames[0]]
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    index_lines = [
        "# Base de conocimiento — reglas de negocio",
        "",
        "Reglas derivadas de `Base de Conocimiento IA VF01.xlsx` (filas ≥79).",
        "Cada regla indica intención, expresiones, cuándo aclarar y respuesta esperada.",
        "",
        "> Ejecutable en runtime: `data/kb_faq_vf01.json` (+ overlays Fase 1).",
        "> Este MD es la memoria humana / RAG-friendly.",
        "",
        "## Índice por producto",
        "",
    ]

    by_prod: dict[str, list[dict]] = {}
    for r in range(from_row, ws.max_row + 1):
        func = _cell(ws, r, 1)
        producto = _cell(ws, r, 5) or "General"
        intencion = _cell(ws, r, 7)
        dato = _cell(ws, r, 8)
        expresiones = _cell(ws, r, 9)
        ambig = _cell(ws, r, 11)
        cuando_aclarar = _cell(ws, r, 12)
        resp_patron = _cell(ws, r, 23)
        resp_esperada = _cell(ws, r, 24)
        if not any([func, expresiones, resp_esperada, resp_patron, dato]):
            continue
        topic = dato or func or f"fila-{r}"
        by_prod.setdefault(producto, []).append(
            {
                "row": r,
                "topic": topic,
                "intent": intencion,
                "expressions": [p.strip() for p in re.split(r"\s*\|\s*", expresiones) if p.strip()],
                "ambiguity": ambig,
                "when_clarify": cuando_aclarar,
                "answer": resp_esperada or resp_patron,
                "func": func,
            }
        )

    for prod in sorted(by_prod.keys()):
        safe = re.sub(r"[^\w\-]+", "_", prod)[:60] or "General"
        path = out_dir / f"{safe}.md"
        items = by_prod[prod]
        index_lines.append(f"- [{prod}]({safe}.md) ({len(items)} reglas)")
        lines = [
            f"# {prod}",
            "",
            f"Fuente: Base de Conocimiento IA VF01 (filas ≥{from_row})",
            "",
        ]
        for e in items:
            lines.append(f"## {e['topic']}")
            lines.append("")
            lines.append(f"- **Fila Excel:** {e['row']}")
            if e["intent"]:
                lines.append(f"- **Intención:** {e['intent']}")
            if e["ambiguity"]:
                lines.append(f"- **¿Ambigüedad?:** {e['ambiguity']}")
            if e["when_clarify"]:
                lines.append(f"- **Cuándo aclarar:** {e['when_clarify']}")
            lines.append("")
            if e["expressions"]:
                lines.append("### Expresiones")
                lines.append("")
                for ex in e["expressions"][:8]:
                    lines.append(f"- {ex}")
                lines.append("")
            if e["answer"]:
                lines.append("### Respuesta esperada")
                lines.append("")
                lines.append(e["answer"])
                lines.append("")
        path.write_text("\n".join(lines), encoding="utf-8")
        written.append(path)

    index_lines.append("")
    idx = out_dir / "README.md"
    idx.write_text("\n".join(index_lines), encoding="utf-8")
    written.append(idx)
    return written


def main() -> int:
    ap = argparse.ArgumentParser(description="Export Excel → docs/ Markdown")
    ap.add_argument(
        "--api-excel",
        default=os.getenv("GENESIS_API_EXCEL_PATH", ""),
        help="Ruta Documentacion_API_Productos.xlsx",
    )
    ap.add_argument(
        "--kb-excel",
        default=os.getenv("GENESIS_KB_EXCEL_PATH", ""),
        help="Ruta Base de Conocimiento IA VF01.xlsx",
    )
    ap.add_argument(
        "--out-dir",
        default=os.getenv("GENESIS_DOCS_OUT_DIR", str(ROOT / "docs")),
    )
    ap.add_argument("--from-row", type=int, default=79)
    args = ap.parse_args()

    out = Path(args.out_dir)
    written_all: list[Path] = []

    if args.api_excel:
        api = Path(args.api_excel)
        if not api.is_file():
            print(f"ERROR: no existe API excel {api}", file=sys.stderr)
            return 1
        written_all += export_api_docs(api, out / "context-api")
        print(f"OK context-api -> {out / 'context-api'} ({len(written_all)} files so far)")
    else:
        print("SKIP context-api (no --api-excel / GENESIS_API_EXCEL_PATH)")

    if args.kb_excel:
        kb = Path(args.kb_excel)
        if not kb.is_file():
            print(f"ERROR: no existe KB excel {kb}", file=sys.stderr)
            return 1
        before = len(written_all)
        written_all += export_kb_rules(kb, out / "business-rules", args.from_row)
        print(f"OK business-rules -> {out / 'business-rules'} (+{len(written_all) - before})")
    else:
        print("SKIP business-rules (no --kb-excel / GENESIS_KB_EXCEL_PATH)")

    print(f"TOTAL written={len(written_all)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
