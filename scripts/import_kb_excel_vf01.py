"""Importa Base de Conocimiento Excel (desde fila N) → FAQ local + markdown KB.

Uso:
  .\\.venv\\Scripts\\python.exe scripts\\import_kb_excel_vf01.py \\
      --excel \"C:\\Users\\...\\Base de Conocimiento IA VF01.xlsx\" \\
      --from-row 79
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: pip install openpyxl"); sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT_JSON = ROOT / "data" / "kb_faq_vf01.json"
DEFAULT_OUT_MD_DIR = ROOT / "Knowledge_Base" / "excel_vf01"


def _norm(text: str) -> str:
    t = (text or "").strip().lower()
    t = t.replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u")
    t = t.replace("¿", "").replace("?", "").replace("¡", "").replace("!", "")
    t = re.sub(r"\s+", " ", t)
    return t


def _split_expressions(raw: str) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"\s*\|\s*", str(raw))
    return [p.strip() for p in parts if p and p.strip()]


def _is_weak_answer(answer: str, topic: str) -> bool:
    a = (answer or "").strip()
    if not a:
        return True
    low = a.lower()
    if low.startswith("tema de información general disponible"):
        return True
    if "basada exclusivamente en el contexto aprobado" in low:
        return True
    if "[respuesta" in low:
        return True
    if a.lower() == (topic or "").strip().lower():
        return True
    return len(a) < 40


def import_rows(excel: Path, from_row: int) -> list[dict]:
    wb = openpyxl.load_workbook(excel, data_only=True)
    ws = wb["Matriz_Cuentas"] if "Matriz_Cuentas" in wb.sheetnames else wb[wb.sheetnames[0]]

    # topic -> best entry
    best_by_topic: dict[str, dict] = {}

    for r in range(from_row, ws.max_row + 1):
        func = ws.cell(r, 1).value
        criterio = ws.cell(r, 2).value or ""
        producto = ws.cell(r, 5).value or ""
        intencion = ws.cell(r, 7).value or ""
        dato = ws.cell(r, 8).value or ""
        expresiones = ws.cell(r, 9).value or ""
        resp_patron = ws.cell(r, 23).value or ""
        resp_esperada = ws.cell(r, 24).value or ""

        if not any([func, criterio, expresiones, resp_esperada, resp_patron]):
            continue

        topic = str(dato or "").strip() or str(func or "").strip()
        candidates = [
            str(resp_esperada or "").strip(),
            str(resp_patron or "").strip(),
        ]
        crit = str(criterio).strip()
        if "\n" in crit:
            candidates.append(crit.split("\n", 1)[1].strip())
        else:
            candidates.append(crit)

        answer = ""
        for cand in candidates:
            if not _is_weak_answer(cand, topic):
                answer = cand
                break
        if not answer:
            continue

        exprs = _split_expressions(str(expresiones))
        if not exprs and topic:
            exprs = [
                f"¿Qué es {topic}?",
                f"Cuéntame sobre {topic}",
                f"¿Qué información tienes sobre {topic}?",
            ]

        entry = {
            "id": f"vf01-r{r}",
            "row": r,
            "topic": topic,
            "product": str(producto).strip(),
            "intent": str(intencion).strip(),
            "expressions": exprs,
            "answer": answer,
            "source": "Base de Conocimiento IA VF01.xlsx",
        }
        key = _norm(topic) or f"row-{r}"
        prev = best_by_topic.get(key)
        if prev is None or len(answer) > len(prev.get("answer") or ""):
            # merge expressions
            if prev:
                merged = list(dict.fromkeys((prev.get("expressions") or []) + exprs))
                entry["expressions"] = merged
            best_by_topic[key] = entry

    return list(best_by_topic.values())


def write_markdown(entries: list[dict], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    # Agrupar por producto
    by_prod: dict[str, list[dict]] = {}
    for e in entries:
        key = e.get("product") or "General"
        by_prod.setdefault(key, []).append(e)

    for prod, items in by_prod.items():
        safe = re.sub(r"[^\w\-]+", "_", prod)[:60] or "General"
        path = out_dir / f"{safe}.md"
        lines = [f"# {prod}", "", "Fuente: Base de Conocimiento IA VF01 (filas ≥79)", ""]
        for e in items:
            lines.append(f"## {e['topic']}")
            lines.append("")
            if e.get("expressions"):
                lines.append("Preguntas ejemplo:")
                for ex in e["expressions"][:4]:
                    lines.append(f"- {ex}")
                lines.append("")
            lines.append(e["answer"])
            lines.append("")
        path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True)
    ap.add_argument("--from-row", type=int, default=79)
    ap.add_argument("--out-json", default=str(DEFAULT_OUT_JSON))
    ap.add_argument("--out-md-dir", default=str(DEFAULT_OUT_MD_DIR))
    args = ap.parse_args()

    excel = Path(args.excel)
    if not excel.is_file():
        print(f"ERROR: no existe {excel}")
        return 1

    entries = import_rows(excel, args.from_row)
    out_json = Path(args.out_json)
    out_json.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": "vf01",
        "from_row": args.from_row,
        "count": len(entries),
        "entries": entries,
    }
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_markdown(entries, Path(args.out_md_dir))
    print(f"OK entries={len(entries)} json={out_json} md={args.out_md_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
