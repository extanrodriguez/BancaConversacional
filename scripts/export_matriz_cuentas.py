"""Export Matriz_Cuentas from Excel to JSON for Test_local reference."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path.home() / "Downloads" / "Consulta de Préstamos.xlsx"
OUT = ROOT / "Test_local" / "data" / "matriz_cuentas.json"


def export(path: Path) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Matriz_Cuentas"]
    headers = [ws.cell(1, c).value for c in range(1, 25)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {headers[i]: ws.cell(r, i + 1).value for i in range(len(headers)) if headers[i]}
        if row.get("Expresiones del cliente"):
            rows.append(row)
    wb.close()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(rows)} rows -> {OUT}")


if __name__ == "__main__":
    export(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX)
